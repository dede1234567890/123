import json
import secrets
import threading
import uuid
from pathlib import Path

from openpyxl import load_workbook

from app import run_sync
from excel_sync import HEADERS
from questions_bank import TOTAL_QUESTIONS
from test_grader import Q9_SOLUTION as Q9, Q10_SOLUTION as Q10, Q11, Q12, Q13, Q14, Q15, Q16, Q17, Q18, Q19, Q20

CORRECT = {
    1: {"answers": {"1": True, "2": True, "3": False, "4": False}},
    2: {"option": "C"},
    3: {"option": "B"},
    4: {"option": "B"},
    5: {"option": "C"},
    6: {"option": "B"},
    7: {"option": "B"},
    8: {"option": "B"},
    9: {"code": Q9},
    10: {"code": Q10},
    11: {"code": Q11},
    12: {"code": Q12},
    13: {"code": Q13},
    14: {"code": Q14},
    15: {"code": Q15},
    16: {"code": Q16},
    17: {"code": Q17},
    18: {"code": Q18},
    19: {"code": Q19},
    20: {"code": Q20},
}

WRONG = {
    1: {"answers": {"1": False, "2": False, "3": True, "4": True}},
    2: {"option": "A"},
    3: {"option": "A"},
    4: {"option": "A"},
    5: {"option": "A"},
    6: {"option": "A"},
    7: {"option": "A"},
    8: {"option": "A"},
    9: {"code": "print('Aprovado')"},
    10: {"code": "aprovados = 3\n"},
    11: {"code": "print('Recusado')"},
    12: {"code": "print(0)\n"},
    13: {"code": "def estoque_baixo(quantidade):\n    return 0\n"},
    14: {"code": "print('Aprovado')"},
    15: {"code": "aprovados = 3\n"},
    16: {"code": "print('Normal')"},
    17: {"code": "def preco_alto(valor):\n    return 0\n"},
    18: {"code": "print(12)\n"},
    19: {"code": "quantidade = 0\n"},
    20: {"code": "grandes = 3\n"},
}


def start_exam(client, nome="Ana Silva", turma="8A", browser=None, token=None):
    browser = browser or str(uuid.uuid4())
    token = token or secrets.token_hex(32)
    response = client.post(
        "/api/start",
        json={"nome": nome, "turma": turma, "browser_session": browser, "token": token},
    )
    return response, browser, token


def auth_headers(attempt_id, token):
    return {"X-Attempt-Id": attempt_id, "Authorization": f"Bearer {token}"}


def finish_exam(client, headers, responses):
    for number in range(1, TOTAL_QUESTIONS + 1):
        state = client.get("/api/state", headers=headers)
        assert state.status_code == 200
        body = state.get_json()
        assert body["question"]["number"] == number
        assert "key" not in body["question"]
        assert "explanation" not in body["question"]
        answered = client.post(
            "/api/answer",
            headers=headers,
            json={"question": number, "response": responses[number]},
        )
        assert answered.status_code == 200, answered.get_json()
        if number < TOTAL_QUESTIONS:
            moved = client.post("/api/next", headers=headers)
            assert moved.status_code == 200
            assert moved.get_json()["question"]["number"] == number + 1
    finished = client.post("/api/finish", headers=headers)
    assert finished.status_code == 200, finished.get_json()
    return finished.get_json()["result"]


def test_identification_rules(client):
    missing = start_exam(client, nome="   ", turma="8A")[0]
    assert missing.status_code == 400
    single = start_exam(client, nome="Ana", turma="8A")[0]
    assert single.status_code == 400
    invalid_class = start_exam(client, nome="Ana Silva", turma="8F")[0]
    assert invalid_class.status_code == 400
    response, _browser, _token = start_exam(client, nome="  Maria   da Silva  ", turma="8C")
    assert response.status_code == 200
    assert response.get_json()["nome"] == "Maria da Silva"
    assert response.get_json()["turma"] == "8C"


def test_perfect_score_is_saved_once_per_class(client, app, sent_messages):
    response, browser, token = start_exam(client, nome="Ana Silva", turma="8A")
    data = response.get_json()
    headers = auth_headers(data["attempt_id"], token)
    result = finish_exam(client, headers, CORRECT)
    assert result["acertos"] == 20
    assert result["erros"] == 0
    assert result["nota"] == 7
    assert result["nota_formatada"] == "7,00"
    assert result["aproveitamento"] == 100
    assert result["salvo"] is True
    assert "foi avisado" in result["registro"]
    assert len(sent_messages) == 1
    assert "Ana Silva" in sent_messages[0].get_content()
    assert "Nota: 7,00" in sent_messages[0].get_content()
    assert "diogo.nascimento@institutojef.org.br" == sent_messages[0]["To"]

    again = client.post("/api/finish", headers=headers)
    assert again.status_code == 200
    assert len(sent_messages) == 1

    restarted, _same_browser, _same_token = start_exam(
        client, nome="Outra Pessoa", turma="8B", browser=browser, token=token
    )
    assert restarted.status_code == 200
    assert restarted.get_json()["finished"] is True
    assert restarted.get_json()["attempt_id"] == data["attempt_id"]
    assert restarted.get_json()["nome"] == "Ana Silva"

    other, _browser_b, token_b = start_exam(client, nome="Bruno Costa", turma="8B")
    other_headers = auth_headers(other.get_json()["attempt_id"], token_b)
    other_result = finish_exam(client, other_headers, WRONG)
    assert other_result["nota"] == 0
    assert other_result["nota_formatada"] == "0,00"
    assert other_result["erros"] == 20

    folder = Path(app.config["DATA_DIR"])
    class_a = load_workbook(folder / "notas_8A.xlsx").active
    class_b = load_workbook(folder / "notas_8B.xlsx").active
    assert [cell.value for cell in class_a[1]] == HEADERS
    assert class_a.max_row == 2
    assert class_a["B2"].value == "Ana Silva"
    assert class_a["H2"].value == 7
    assert class_b.max_row == 2
    assert class_b["B2"].value == "Bruno Costa"
    assert class_b["H2"].value == 0
    assert all(cell.value != "Bruno Costa" for cell in class_a["B"])


def test_answers_cannot_be_changed_or_skipped(client):
    response, _browser, token = start_exam(client)
    headers = auth_headers(response.get_json()["attempt_id"], token)
    first = client.post("/api/answer", headers=headers, json={"question": 1, "response": CORRECT[1]})
    assert first.status_code == 200
    assert first.get_json()["correct"] is True
    assert first.get_json()["points_formatada"] == "0,20"
    repeated = client.post("/api/answer", headers=headers, json={"question": 1, "response": WRONG[1]})
    assert repeated.status_code == 409
    skipped = client.post("/api/answer", headers=headers, json={"question": 2, "response": CORRECT[2]})
    assert skipped.status_code == 409
    state = client.get("/api/state", headers=headers)
    assert state.get_json()["progress"]["acertos"] == 1
    assert state.get_json()["progress"]["pontuacao"] == 0.2
    assert state.get_json()["progress"]["pontuacao_formatada"] == "0,20"
    assert state.get_json()["feedback"]["points_formatada"] == "0,20"
    assert "token_hash" not in json.dumps(state.get_json())


def test_parallel_confirmations_do_not_duplicate_points(app, client):
    response, _browser, token = start_exam(client)
    attempt_id = response.get_json()["attempt_id"]
    store = app.extensions["store"]
    barrier = threading.Barrier(2)

    def confirm():
        barrier.wait()
        return store.save_answer(attempt_id, 1, '{"answers": {}}', True)

    with_threads = []
    first = threading.Thread(target=lambda: with_threads.append(confirm()))
    second = threading.Thread(target=lambda: with_threads.append(confirm()))
    first.start()
    second.start()
    first.join()
    second.join()
    conflicts = [item["conflict"] for item in with_threads]
    assert conflicts.count(False) == 1
    assert conflicts.count(True) == 1
    saved = store.get_attempt(attempt_id)
    assert saved["acertos"] == 1
    assert len(store.list_answers(attempt_id)) == 1


def test_failed_email_can_be_retried_without_a_new_grade(client, app, sent_messages):
    mailer = app.extensions["mailer"]
    original = mailer._sender

    def fail_once(*args):
        mailer._sender = original
        raise OSError("smtp indisponível")

    mailer._sender = fail_once
    response, _browser, token = start_exam(client, nome="Carla Souza", turma="8D")
    headers = auth_headers(response.get_json()["attempt_id"], token)
    result = finish_exam(client, headers, CORRECT)
    assert result["salvo"] is True
    assert "reenviado automaticamente" in result["registro"]
    assert sent_messages == []

    retried = client.post("/api/finish", headers=headers)
    assert retried.status_code == 200
    assert "foi avisado" in retried.get_json()["result"]["registro"]
    assert len(sent_messages) == 1
    folder = Path(app.config["DATA_DIR"])
    sheet = load_workbook(folder / "notas_8D.xlsx").active
    assert sheet.max_row == 2
    assert sheet["J2"].value == "Salvo e sincronizado"


def test_admin_export_requires_login_and_separates_classes(client, app):
    denied = client.get("/api/admin/results")
    assert denied.status_code == 401
    bad_login = client.post("/api/admin/login", json={"username": "diogo", "password": "errada"})
    assert bad_login.status_code == 401
    wrong_user = client.post("/api/admin/login", json={"username": "aluno", "password": "segredo-teste"})
    assert wrong_user.status_code == 401
    login = client.post("/api/admin/login", json={"username": "diogo", "password": "segredo-teste"})
    assert login.status_code == 200
    response, _browser, token = start_exam(client, nome="Diego Lima", turma="8E")
    finish_exam(client, auth_headers(response.get_json()["attempt_id"], token), CORRECT)
    exported = client.get("/api/admin/export/8E")
    assert exported.status_code == 200
    assert "spreadsheetml" in exported.content_type
    listing = client.get("/api/admin/results?turma=8E")
    rows = listing.get_json()["results"]
    assert len(rows) == 1
    assert rows[0]["nome"] == "Diego Lima"
    assert rows[0]["nota"] == 7
    assert rows[0]["nota_formatada"] == "7,00"
    assert "token_hash" not in listing.get_data(as_text=True)
    missing = client.get("/backend/questions_bank.py")
    assert missing.status_code == 404
    script = client.get("/js/app.js")
    assert script.status_code == 200
    assert b"explanation_correct" not in script.data
    run_sync(app)


def test_resume_keeps_the_same_attempt(client):
    response, browser, token = start_exam(client, nome="Eva Martins", turma="8A")
    attempt_id = response.get_json()["attempt_id"]
    headers = auth_headers(attempt_id, token)
    client.post("/api/answer", headers=headers, json={"question": 1, "response": WRONG[1]})
    resumed, _browser, _token = start_exam(
        client, nome="Nome Trocado", turma="8E", browser=browser, token=token
    )
    assert resumed.get_json()["attempt_id"] == attempt_id
    assert resumed.get_json()["nome"] == "Eva Martins"
    assert resumed.get_json()["turma"] == "8A"
    state = client.get("/api/state", headers=headers)
    assert state.get_json()["progress"]["erros"] == 1
    foreign = start_exam(client, nome="Eva Martins", turma="8A", browser=browser, token=secrets.token_hex(32))[0]
    assert foreign.status_code == 409


def test_security_closure_is_written_to_the_spreadsheet(client, app, sent_messages):
    response, _browser, token = start_exam(client, nome="Caio Mendes", turma="8C")
    headers = auth_headers(response.get_json()["attempt_id"], token)
    denied = client.post("/api/finish", headers=headers)
    assert denied.status_code == 409
    closed = client.post("/api/finish", headers=headers, json={"encerramento": "seguranca"})
    assert closed.status_code == 200
    result = closed.get_json()["result"]
    assert result["acertos"] == 0
    assert "capturar a tela" in result["mensagem"]
    assert result["encerramento"] == "Prova encerrada por tentativa de burlar a segurança"
    sheet = load_workbook(Path(app.config["DATA_DIR"]) / "notas_8C.xlsx").active
    assert sheet["B2"].value == "Caio Mendes"
    assert sheet["K2"].value == "Prova encerrada por tentativa de burlar a segurança"
    assert "burlar a segurança" in sent_messages[0].get_content()
