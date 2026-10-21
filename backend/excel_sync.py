"""Gera as cinco planilhas a partir do armazenamento central."""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from questions_bank import MAX_SCORE_CENTS
from storage import TURMAS, format_br

HEADERS = [
    "Identificador da tentativa",
    "Nome completo",
    "Turma",
    "Início",
    "Término",
    "Acertos",
    "Erros",
    "Nota final",
    "Aproveitamento (%)",
    "Situação do salvamento",
    "Encerramento",
]


def encerramento_texto(row: dict) -> str:
    if row.get("encerramento") == "seguranca":
        return "Prova encerrada por tentativa de burlar a segurança"
    return "Concluída pelo aluno"


def situacao(row: dict) -> str:
    if not row.get("saved"):
        return "Não salvo"
    pendencias = []
    if not row.get("excel_synced"):
        pendencias.append("planilha pendente")
    if not row.get("email_sent"):
        pendencias.append("e-mail pendente")
    if not pendencias:
        return "Salvo e sincronizado"
    return "Salvo (" + " e ".join(pendencias) + ")"


def excel_text(value: str) -> str:
    if value and value[0] in "=+-@":
        return "'" + value
    return value


def workbook_path(data_dir: Path, turma: str) -> Path:
    if turma not in TURMAS:
        raise ValueError("Turma inválida.")
    return Path(data_dir) / f"notas_{turma}.xlsx"


def write_workbook(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Notas"
    sheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="0F6B56")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri")
    thin = Border(
        left=Side(style="thin", color="D5E2DB"),
        right=Side(style="thin", color="D5E2DB"),
        top=Side(style="thin", color="D5E2DB"),
        bottom=Side(style="thin", color="D5E2DB"),
    )
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        acertos = int(row["acertos"])
        cents = int(row.get("nota_cents") or 0)
        sheet.append(
            [
                excel_text(row["attempt_id"]),
                excel_text(row["nome"]),
                row["turma"],
                format_br(row["started_at"]),
                format_br(row["finished_at"]),
                acertos,
                int(row["erros"]),
                round(cents / 100, 2),
                int(cents * 100 / MAX_SCORE_CENTS),
                situacao(row),
                encerramento_texto(row),
            ]
        )
    for data_row in sheet.iter_rows(min_row=2):
        for cell in data_row:
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
        data_row[7].number_format = "0.00"
        data_row[8].number_format = "0"
    widths = [40, 36, 12, 20, 20, 12, 12, 14, 22, 36, 62]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = f"A1:K{max(1, sheet.max_row)}"
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 28
    temporary = path.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def sync_class(store, data_dir: Path, turma: str) -> Path:
    rows = store.list_finished(turma)
    path = workbook_path(data_dir, turma)
    # A situação "planilha pendente" some depois que a gravação termina.
    prepared = []
    for row in rows:
        copied = dict(row)
        copied["excel_synced"] = True
        prepared.append(copied)
    write_workbook(path, prepared)
    store.mark_excel_synced([row["attempt_id"] for row in rows])
    return path


def sync_all(store, data_dir: Path) -> list[str]:
    written = []
    for turma in TURMAS:
        path = sync_class(store, data_dir, turma)
        written.append(str(path))
    return written
